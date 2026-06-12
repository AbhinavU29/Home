from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from app.core.database import get_db
from app.api.auth import get_current_user, RoleChecker, validate_password_strength
from app.core.security import get_password_hash
from app.models.models import User, Product, AuditLog
import openpyxl
import io
import re

router = APIRouter(prefix="/import", tags=["import-data"])

def read_excel_rows(file_bytes: bytes, filename: str) -> list:
    rows = []
    if filename.endswith(".xlsx"):
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        sheet = wb.active
        for r in sheet.iter_rows(values_only=True):
            # Convert tuples to lists
            rows.append(list(r))
    elif filename.endswith(".xls"):
        try:
            import xlrd
        except ImportError:
            raise HTTPException(
                status_code=400,
                detail="The library 'xlrd' is required to parse .xls files. Please contact administrator or upload as .xlsx."
            )
        wb = xlrd.open_workbook(file_contents=file_bytes)
        sheet = wb.sheet_by_index(0)
        for r_idx in range(sheet.nrows):
            rows.append([sheet.cell_value(r_idx, c_idx) for c_idx in range(sheet.ncols)])
    else:
        raise HTTPException(
            status_code=400,
            detail="Unsupported file format. Only Excel .xlsx and .xls formats are supported."
        )
    return rows

@router.post("/users")
def import_users(
    file: UploadFile = File(...),
    dry_run: bool = False,
    db: Session = Depends(get_db),
    current_user: User = Depends(RoleChecker(["Admin"]))
):
    contents = file.file.read()
    try:
        rows = read_excel_rows(contents, file.filename)
    except Exception as e:
        if isinstance(e, HTTPException):
            raise e
        raise HTTPException(status_code=400, detail=f"Failed to read Excel file: {str(e)}")

    if not rows or len(rows) < 2:
        raise HTTPException(status_code=400, detail="The uploaded Excel sheet is empty or contains no data rows.")

    headers = [str(h).strip().lower() for h in rows[0] if h is not None]
    
    # Map headers dynamically
    name_idx = next((i for i, h in enumerate(headers) if "name" in h and "user" not in h and "phone" not in h and "mobile" not in h), -1)
    username_idx = next((i for i, h in enumerate(headers) if "username" in h or "user name" in h), -1)
    password_idx = next((i for i, h in enumerate(headers) if "password" in h or "pass" in h), -1)
    role_idx = next((i for i, h in enumerate(headers) if "role" in h), -1)
    mobile_idx = next((i for i, h in enumerate(headers) if "mobile" in h or "phone" in h or "contact" in h), -1)

    if username_idx == -1 or password_idx == -1:
        raise HTTPException(
            status_code=400,
            detail="Required headers missing. Make sure your sheet contains 'Username' and 'Password' columns."
        )

    success_count = 0
    errors = []
    processed_usernames = set()

    for idx, row in enumerate(rows[1:], 2):
        # Pad row with None values if shorter than headers
        while len(row) < len(headers):
            row.append(None)
            
        username = str(row[username_idx]).strip() if row[username_idx] is not None else ""
        password = str(row[password_idx]).strip() if row[password_idx] is not None else ""
        name = str(row[name_idx]).strip() if (name_idx != -1 and row[name_idx] is not None) else None
        role = str(row[role_idx]).strip() if (role_idx != -1 and row[role_idx] is not None) else "Operator"
        mobile = str(row[mobile_idx]).strip() if (mobile_idx != -1 and row[mobile_idx] is not None) else None

        if not username:
            errors.append({"row": idx, "error": "Username is empty."})
            continue
        if not password:
            errors.append({"row": idx, "error": "Password is empty."})
            continue

        # Check for duplicates inside the uploaded file
        if username.lower() in processed_usernames:
            errors.append({"row": idx, "error": f"Duplicate username '{username}' found inside this import file."})
            continue
        processed_usernames.add(username.lower())

        # Check for duplicates in the database
        db_exists = db.query(User).filter(User.username == username).first()
        if db_exists:
            errors.append({"row": idx, "error": f"Username '{username}' already exists in database."})
            continue

        # Password strength validation
        try:
            validate_password_strength(password)
        except HTTPException as strength_err:
            errors.append({"row": idx, "error": f"Password: {strength_err.detail}"})
            continue

        # Normalize role
        normalized_role = "Operator"
        if role.lower() in ["admin", "administrator"]:
            normalized_role = "Admin"
        elif role.lower() in ["operator", "cashier", "billing", "staff"]:
            normalized_role = "Operator"

        if not dry_run:
            try:
                new_user = User(
                    username=username,
                    password=get_password_hash(password),
                    role=normalized_role,
                    full_name=name,
                    mobile_number=mobile,
                    is_active=1
                )
                db.add(new_user)
                db.flush()
                success_count += 1
            except Exception as insert_err:
                db.rollback()
                errors.append({"row": idx, "error": f"Database error: {str(insert_err)}"})
        else:
            success_count += 1

    if not dry_run:
        db.commit()

        # Create Audit Log for Import
        if success_count > 0:
            audit = AuditLog(
                action="IMPORT_USERS_EXCEL",
                table_name="users",
                record_id=0,
                old_values=None,
                new_values=f"Imported {success_count} users successfully from excel sheet.",
                user_id=current_user.user_id
            )
            db.add(audit)
            db.commit()

    return {
        "total_rows": len(rows) - 1,
        "success_count": success_count,
        "failure_count": len(errors),
        "errors": errors
    }

@router.post("/products")
def import_products(
    file: UploadFile = File(...),
    dry_run: bool = False,
    db: Session = Depends(get_db),
    current_user: User = Depends(RoleChecker(["Admin"]))
):
    contents = file.file.read()
    try:
        rows = read_excel_rows(contents, file.filename)
    except Exception as e:
        if isinstance(e, HTTPException):
            raise e
        raise HTTPException(status_code=400, detail=f"Failed to read Excel file: {str(e)}")

    if not rows or len(rows) < 2:
        raise HTTPException(status_code=400, detail="The uploaded Excel sheet is empty or contains no data rows.")

    headers = [str(h).strip().lower() for h in rows[0] if h is not None]

    # Map headers dynamically
    name_idx = next((i for i, h in enumerate(headers) if "product" in h or "name" in h), -1)
    type_idx = next((i for i, h in enumerate(headers) if "ornament" in h or "type" in h), -1)
    weight_idx = next((i for i, h in enumerate(headers) if "weight" in h or "wt" in h), -1)
    charge_idx = next((i for i, h in enumerate(headers) if "charge" in h or "making" in h), -1)
    gst_idx = next((i for i, h in enumerate(headers) if "gst" in h or "tax" in h), -1)

    if name_idx == -1 or type_idx == -1:
        raise HTTPException(
            status_code=400,
            detail="Required headers missing. Make sure your sheet contains 'Product Name' and 'Ornament Type' columns."
        )

    success_count = 0
    errors = []
    processed_combos = set()

    for idx, row in enumerate(rows[1:], 2):
        while len(row) < len(headers):
            row.append(None)

        name = str(row[name_idx]).strip() if row[name_idx] is not None else ""
        ornament = str(row[type_idx]).strip() if row[type_idx] is not None else ""
        
        weight_raw = row[weight_idx] if (weight_idx != -1 and row[weight_idx] is not None) else 0.0
        charge_raw = row[charge_idx] if (charge_idx != -1 and row[charge_idx] is not None) else 0.0
        gst_raw = row[gst_idx] if (gst_idx != -1 and row[gst_idx] is not None) else 3.0

        if not name:
            errors.append({"row": idx, "error": "Product Name is empty."})
            continue
        if not ornament:
            errors.append({"row": idx, "error": "Ornament Type is empty."})
            continue

        # Try to parse numeric fields
        try:
            weight = float(weight_raw)
            if weight < 0:
                raise ValueError
        except ValueError:
            errors.append({"row": idx, "error": f"Invalid Weight value '{weight_raw}'. Must be a non-negative number."})
            continue

        try:
            charge = float(charge_raw)
            if charge < 0:
                raise ValueError
        except ValueError:
            errors.append({"row": idx, "error": f"Invalid Making Charge value '{charge_raw}'. Must be a non-negative number."})
            continue

        try:
            gst = float(gst_raw)
            if gst < 0:
                raise ValueError
        except ValueError:
            errors.append({"row": idx, "error": f"Invalid GST percentage value '{gst_raw}'. Must be a non-negative number."})
            continue

        # Duplicate checking inside the upload file
        combo = (name.lower(), ornament.lower())
        if combo in processed_combos:
            errors.append({"row": idx, "error": f"Duplicate product '{name}' of type '{ornament}' inside this import file."})
            continue
        processed_combos.add(combo)

        # Duplicate checking in database
        db_exists = db.query(Product).filter(
            Product.product_name == name,
            Product.ornament_type == ornament
        ).first()
        if db_exists:
            errors.append({"row": idx, "error": f"Product '{name}' with type '{ornament}' already exists in database."})
            continue

        if not dry_run:
            try:
                new_prod = Product(
                    product_name=name,
                    ornament_type=ornament,
                    default_gold_weight=weight,
                    default_making_charge=charge,
                    gst_percent=gst,
                    status="Active"
                )
                db.add(new_prod)
                db.flush()
                success_count += 1
            except Exception as insert_err:
                db.rollback()
                errors.append({"row": idx, "error": f"Database error: {str(insert_err)}"})
        else:
            success_count += 1

    if not dry_run:
        db.commit()

        if success_count > 0:
            audit = AuditLog(
                action="IMPORT_PRODUCTS_EXCEL",
                table_name="products",
                record_id=0,
                old_values=None,
                new_values=f"Imported {success_count} products successfully from excel sheet.",
                user_id=current_user.user_id
            )
            db.add(audit)
            db.commit()

    return {
        "total_rows": len(rows) - 1,
        "success_count": success_count,
        "failure_count": len(errors),
        "errors": errors
    }

@router.get("/users/template")
def get_users_template(current_user: User = Depends(RoleChecker(["Admin"]))):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Users Template"
    headers = ["Name", "Username", "Password", "Role", "Mobile"]
    ws.append(headers)
    ws.append(["John Doe", "john_operator", "OperatorPass123!", "Operator", "9876543210"])
    ws.append(["Jane Admin", "jane_admin", "AdminPass123!", "Admin", "9876543211"])
    
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    
    return StreamingResponse(
        out,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=users_import_template.xlsx"}
    )

@router.get("/products/template")
def get_products_template(current_user: User = Depends(RoleChecker(["Admin"]))):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Products Template"
    headers = ["Product Name", "Ornament Type", "Gold Weight", "Making Charge", "GST"]
    ws.append(headers)
    ws.append(["Gold Ring 22K", "Ring", "4.5", "250", "3"])
    ws.append(["Gold Choker BIS", "Necklace", "18.2", "450", "3"])
    
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    
    return StreamingResponse(
        out,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=products_import_template.xlsx"}
    )
