from fastapi import APIRouter, Depends, Response
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import func
from app.core.database import get_db
from app.api.auth import get_current_user
from app.models.models import Bill, Customer, User, BillItem
from app.schemas.schemas import DashboardKPI
from typing import Optional, List
from datetime import datetime, date
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

router = APIRouter(prefix="/reports", tags=["reports"])

@router.get("/dashboard-kpi", response_model=DashboardKPI)
def get_dashboard_summary(db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    # Total sales sum
    total_sales = db.query(func.sum(Bill.total_amount)).scalar() or 0.0
    
    # Invoices count
    bills_count = db.query(func.count(Bill.bill_id)).scalar() or 0
    
    # Customer count
    customers_count = db.query(func.count(Customer.customer_id)).scalar() or 0
    
    # Last 5 bills
    recent_bills = db.query(Bill).order_by(Bill.date.desc(), Bill.bill_id.desc()).limit(5).all()
    
    return {
        "total_sales": total_sales,
        "bills_count": bills_count,
        "customers_count": customers_count,
        "recent_bills": recent_bills
    }

def parse_date_start(date_str: str):
    if not date_str:
        return None
    try:
        return datetime.strptime(date_str, "%Y-%m-%d")
    except ValueError:
        try:
            return datetime.strptime(date_str, "%Y-%m-%d %H:%M:%S")
        except ValueError:
            try:
                return datetime.strptime(date_str.replace("T", " "), "%Y-%m-%d %H:%M")
            except ValueError:
                return None

def parse_date_end(date_str: str):
    if not date_str:
        return None
    try:
        d = datetime.strptime(date_str, "%Y-%m-%d")
        return d.replace(hour=23, minute=59, second=59)
    except ValueError:
        try:
            return datetime.strptime(date_str, "%Y-%m-%d %H:%M:%S")
        except ValueError:
            try:
                return datetime.strptime(date_str.replace("T", " "), "%Y-%m-%d %H:%M")
            except ValueError:
                return None

@router.get("/export/excel")
def export_excel(
    from_date: Optional[str] = None,
    to_date: Optional[str] = None,
    purchase_id: Optional[str] = None,
    customer_name: Optional[str] = None,
    phone_number: Optional[str] = None,
    product_name: Optional[str] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    """
    Generates and exports an Excel sales report matching the selected query filters.
    Includes itemized details for each bill.
    """
    query = db.query(Bill)
    
    if customer_name:
        query = query.filter(func.lower(Bill.customer_name).like(f"%{customer_name.lower()}%"))
    if phone_number:
        query = query.filter(Bill.phone_number.like(f"%{phone_number}%"))
    if purchase_id:
        query = query.filter(func.lower(Bill.purchase_id).like(f"%{purchase_id.lower()}%"))
    if product_name:
        query = query.join(Bill.items).filter(func.lower(BillItem.product_name).like(f"%{product_name.lower()}%")).distinct()
    if from_date:
        fd = parse_date_start(from_date)
        if fd:
            query = query.filter(Bill.date >= fd)
    if to_date:
        td = parse_date_end(to_date)
        if td:
            query = query.filter(Bill.date <= td)
        
    bills = query.order_by(Bill.date.desc(), Bill.bill_id.desc()).all()
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales Summary Report"
    
    # Enable grid lines visibility in Excel
    ws.views.sheetView[0].showGridLines = True
    
    headers = [
        "Purchase ID", "Customer Name", "Date", "Created By", "Product Name", "Ornament Type",
        "Quantity", "Amount", "GST", "Total Sales"
    ]
    ws.append(headers)
    
    # Premium theme styling (Navy headers + Calibri typography)
    header_fill = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    center_align = Alignment(horizontal="center", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    
    row_num = 2
    for b in bills:
        for item in b.items:
            # Base amount calculation
            bullion = item.weight_gold_g * item.rate_per_gram
            mrp_val = item.quantity * item.MRP
            item_sub = bullion + mrp_val + item.making_charge
            item_discount = getattr(item, 'discount', 0.0)
            
            amount = max(0.0, item_sub - item_discount)
            gst_amount = amount * (item.GST / 100.0)
            total_sales = amount + gst_amount
            
            creator_details = f"{b.creator.username} ({b.creator.full_name or 'N/A'})" if b.creator else str(b.created_by)
            
            ws.append([
                b.purchase_id,
                b.customer_name,
                b.date.strftime("%d-%m-%Y %I:%M %p (IST)") if isinstance(b.date, (date, datetime)) else str(b.date),
                creator_details,
                item.product_name,
                item.ornament_type or "",
                item.quantity,
                amount,
                gst_amount,
                total_sales
            ])
            
            # Formatting and alignments
            for c_idx in range(1, 11):
                cell = ws.cell(row=row_num, column=c_idx)
                cell.border = thin_border
                cell.font = Font(name="Calibri", size=10)
                if c_idx in [1, 2, 4, 5, 6]:
                    cell.alignment = left_align
                elif c_idx in [3, 7]:
                    cell.alignment = center_align
                elif c_idx in [8, 9, 10]:
                    cell.alignment = right_align
                    # Standard Currency format
                    cell.number_format = '"₹"#,##0.00'
            row_num += 1
            
    # Auto-adjust column widths dynamically
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 10)
        
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    
    headers = {
        'Content-Disposition': 'attachment; filename="Prasanth_Jewellery_Sales_Report.xlsx"'
    }
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers
    )
